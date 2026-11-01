from app.pipeline import content, package


def test_forbidden_hits_and_blocks():
    assert content.forbidden_hits("최고의 텀블러", "항균 99% 효과") == ["최고", "항균 99"]
    assert content.forbidden_hits("스테인리스 텀블러 500ml") == []
    b = content.blocks_for("CN", 12)
    assert "중국" in b["notice"] and "12일" in b["notice"] and "개인통관고유부호" in b["notice"]
    assert "150달러" in b["customs"] and "반송" in b["returns"]
    custom = content.blocks_for("JP", 8, "{country} 발송, {lead_days}일")
    assert custom["notice"] == "일본 발송, 8일"


def test_description_html_escapes_and_includes_blocks():
    out = content.ContentOut(title_options=["a"], bullet_points=["<b>굵게</b>", "두번째"], description_paragraphs=["문단 & 기호"], options_ko=[], notes="")
    html = content.description_html(out, content.blocks_for("CN", 12))
    assert "&lt;b&gt;" in html and "문단 &amp; 기호" in html and "<h4>구매대행 안내</h4>" in html


def test_card2_text():
    cand = {"suggested_price_krw": 15400, "margin_rate": 0.66, "cost_breakdown_json": {"total_cost_krw": 12309}}
    cont = {"title_chosen": "보온 텀블러 360ml 손잡이", "title_options": ["보온 텀블러 360ml 손잡이", "학생 보온병 360ml"],
            "options_ko_json": [{"name_ko": "블루 / 200ml", "sell": True}, {"name_ko": "핑크 / 360ml", "sell": False}], "image_urls": ["u"], "mode": "A", "forbidden_hits": ["최고"]}
    t = content.card2_text(cand, cont, {"url": "https://x"})
    assert "15,400원" in t and "옵션 2개 (판매 1)" in t and "금지 문구 감지  최고" in t and "https://x" in t


def test_excel_row_and_append(tmp_path):
    cand = {"id": 17, "suggested_price_krw": 15400}
    cont = {"title_chosen": "보온 텀블러 360ml", "description_html": "<p>x</p>", "options_ko_json": [{"name_ko": "블루", "sell": True}, {"name_ko": "핑크", "sell": False}]}
    item = {"url": "https://detail.1688.com/offer/1.html", "source_item_id": "1"}
    row = package.excel_row(cand, cont, item, ["https://p/1.jpg", "https://p/2.jpg"], naver_cat_id="50004540", country="CN")
    assert row["상품명"] == "보온 텀블러 360ml" and row["판매가"] == 15400 and row["대표이미지"].endswith("1.jpg") and row["추가이미지"].endswith("2.jpg")
    assert row["옵션값"] == "블루" and row["원산지"] == "중국" and row["카테고리ID"] == "50004540"
    xlsx = tmp_path / "out" / "naver_bulk_upload.xlsx"
    package.append_excel(xlsx, row)
    package.append_excel(xlsx, row)
    from openpyxl import load_workbook

    ws = load_workbook(xlsx).active
    assert ws.max_row == 3 and ws.cell(1, 1).value == "상품명" and ws.cell(2, 2).value == 15400
