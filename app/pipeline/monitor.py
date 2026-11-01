"""⑩ 등록 상품 감시 (DESIGN.md 2장 "등록 후" · 6-⑩). listed · suspended 상품의 소스를 재수집해 변동을 알린다.

  python -m app.pipeline.monitor                 # 내 listings 전부 (요청 정책은 소스 설정 공유)
  python -m app.pipeline.monitor --cand 17
  python -m app.pipeline.monitor --mark 17 --store-no 1234567 --price 15400   # 관리자가 "등록 완료" 표시 (packaged → listed)

변동 판정: 최소 수량 구간 단가 ±임계값(기본 5%) · 최소 수량 변경 · 전체 품절 · 등록 옵션 품절 · 페이지 삭제 · 복구.
같은 상품의 같은 변동은 결정(판매가 수정 / 판매 중지 / 무시) 전까지 다시 알리지 않는다. 3회 연속 실패는 check_needed.
"""
from __future__ import annotations

import argparse
import sys
from datetime import date, datetime, timezone
from pathlib import Path

from app import repo
from app.config import ROOT, settings


def _now():
    return datetime.now(timezone.utc).isoformat()


# ───────────────────────── 순수 로직 ─────────────────────────
def snapshot(price_tiers: list[dict], moq: int, options: list[dict] | None, sold_out: bool, *, page_missing: bool = False) -> dict:
    from app.pipeline.scoring import unit_price_at_moq

    unit, m = unit_price_at_moq(price_tiers or [], moq or 1)
    opts = options or []
    return {
        "unit_price": unit, "moq": m, "sold_out": bool(sold_out), "page_missing": page_missing,
        "options_soldout": sorted(o.get("name_raw") or o.get("name") or "" for o in opts if o.get("sold_out")),
        "option_count": len(opts),
    }


def diff(before: dict, after: dict, *, price_threshold: float = 0.05) -> list[dict]:
    """변동 목록 [{kind, before, after}]. 임계값 미만 가격 변동은 기록하지 않는다(알림도 없음)."""
    ev: list[dict] = []
    if after.get("page_missing"):
        if not before.get("page_missing"):
            ev.append({"kind": "removed", "before": {"page_missing": False}, "after": {"page_missing": True}})
        return ev
    if before.get("page_missing") or (before.get("sold_out") and not after.get("sold_out")):
        ev.append({"kind": "restored", "before": {"sold_out": before.get("sold_out"), "page_missing": before.get("page_missing", False)}, "after": {"sold_out": False}})
    if after.get("sold_out") and not before.get("sold_out"):
        ev.append({"kind": "soldout", "before": {"sold_out": False}, "after": {"sold_out": True}})
    bp, ap = float(before.get("unit_price") or 0), float(after.get("unit_price") or 0)
    if bp > 0 and ap > 0 and abs(ap - bp) / bp >= price_threshold:
        ev.append({"kind": "price", "before": {"unit_price": bp}, "after": {"unit_price": ap, "change_pct": round((ap - bp) / bp * 100, 1)}})
    if int(before.get("moq") or 1) != int(after.get("moq") or 1):
        ev.append({"kind": "moq", "before": {"moq": before.get("moq")}, "after": {"moq": after.get("moq")}})
    newly = sorted(set(after.get("options_soldout") or []) - set(before.get("options_soldout") or []))
    if newly and not after.get("sold_out"):
        ev.append({"kind": "option_soldout", "before": {"options_soldout": before.get("options_soldout") or []}, "after": {"options_soldout": after.get("options_soldout") or [], "newly": newly}})
    return ev


def eta(source_ordered: date | None, lead_days: int) -> date | None:
    from datetime import timedelta

    return (source_ordered + timedelta(days=int(lead_days))) if source_ordered else None


# ───────────────────────── 등록 완료 표시 ─────────────────────────
def mark_listed(c, uid: str, cand_id: int, *, store_product_no: str | None, sale_price_krw: int) -> dict:
    cand = c.table("candidates").select("*").eq("id", cand_id).limit(1).execute().data[0]
    item = c.table("source_items").select("*").eq("id", cand["source_item_id"]).limit(1).execute().data[0]
    det = (c.table("source_item_details").select("options_json").eq("source_item_id", item["id"]).limit(1).execute().data or [{}])[0]
    base = snapshot(item.get("price_tiers_json") or [], item.get("moq") or 1, det.get("options_json") or [], bool(item.get("sold_out")))
    row = {"candidate_id": cand_id, "user_id": uid, "source_item_id": item["id"], "store_product_no": store_product_no, "sale_price_krw": int(sale_price_krw),
           "listed_at": _now(), "suspended_at": None, "monitor_status": "ok", "baseline_json": base}
    c.table("listings").upsert(row, on_conflict="candidate_id").execute()
    c.table("candidates").update({"status": "listed"}).eq("id", cand_id).execute()
    return row


# ───────────────────────── 재점검 ─────────────────────────
def check_one(c, uid: str, listing: dict, *, headed: bool = False) -> list[dict]:
    from app.pipeline.collector import make_adapter
    from app.pipeline.scoring import compute_economics
    from app.sources.base import Blocked

    item = c.table("source_items").select("*").eq("id", listing["source_item_id"]).limit(1).execute().data[0]
    src = c.table("sources").select("*").eq("id", item["source_id"]).limit(1).execute().data[0]
    us = repo.my_user_sources(c).get(item["source_id"]) or {}
    try:
        with make_adapter(src["adapter_key"], us, headed=headed) as ad:
            d = ad.detail(item["url"])
        after = snapshot(d.price_tiers, d.moq, d.options, d.sold_out, page_missing=("商品不存在" in (d.title_raw or "")) or (not d.price_tiers and d.sold_out))
        # 공용 사실 갱신 (모두에게 최신 가격)
        c.table("source_items").update({"price_tiers_json": d.price_tiers, "moq": d.moq, "sold_out": d.sold_out, "last_price_src": (d.price_tiers[0]["unit_price"] if d.price_tiers else None), "last_seen_at": _now()}).eq("id", item["id"]).execute()
        if d.options:
            c.table("source_item_details").upsert({"source_item_id": item["id"], "options_json": d.options, "fetched_at": _now()}, on_conflict="source_item_id").execute()
    except Blocked as e:
        n = int(listing.get("consecutive_failures") or 0) + 1
        c.table("listings").update({"consecutive_failures": n, "last_checked_at": _now(), "monitor_status": "check_needed" if n >= 3 else listing["monitor_status"]}).eq("candidate_id", listing["candidate_id"]).execute()
        return [{"error": f"차단: {e}"}]
    except Exception as e:
        n = int(listing.get("consecutive_failures") or 0) + 1
        c.table("listings").update({"consecutive_failures": n, "last_checked_at": _now(), "monitor_status": "check_needed" if n >= 3 else listing["monitor_status"]}).eq("candidate_id", listing["candidate_id"]).execute()
        return [{"error": f"{type(e).__name__}: {e}"}]

    events = diff(listing.get("baseline_json") or {}, after, price_threshold=settings.monitor_price_threshold)
    open_kinds = {r["kind"] for r in c.table("monitor_events").select("kind").eq("candidate_id", listing["candidate_id"]).is_("decision", "null").execute().data}
    created = []
    for ev in events:
        if ev["kind"] in open_kinds:
            continue  # 결정 전까지 다시 알리지 않음
        margin = suggested = None
        if ev["kind"] in ("price", "moq"):
            sig = (c.table("market_signals").select("*").eq("source_item_id", item["id"]).limit(1).execute().data or [None])[0]
            new_item = {**item, "price_tiers_json": d.price_tiers, "moq": d.moq}
            det = (c.table("source_item_details").select("package_weight_g").eq("source_item_id", item["id"]).limit(1).execute().data or [{}])[0]
            econ = compute_economics(new_item, sig, _rates(c), weight_g=float(det.get("package_weight_g") or 0))
            # 현재 판매가 기준 재계산 마진
            sp = float(listing["sale_price_krw"])
            margin = round((sp - econ.total_cost_krw) / sp, 4) if sp else None
            suggested = econ.suggested_price_krw
        row = {"candidate_id": listing["candidate_id"], "user_id": uid, "kind": ev["kind"], "before_json": ev["before"], "after_json": ev["after"],
               "recalculated_margin": margin, "suggested_price_krw": suggested}
        created.append(c.table("monitor_events").insert(row).execute().data[0])
    status = "removed" if after.get("page_missing") else "soldout" if after.get("sold_out") else ("changed" if created else "ok")
    c.table("listings").update({"last_checked_at": _now(), "consecutive_failures": 0, "monitor_status": status}).eq("candidate_id", listing["candidate_id"]).execute()
    return created


def _rates(c):
    from app.adapters.exchange_rate import get_rates

    return get_rates(client=c)


KIND_KO = {"price": "가격 변동", "soldout": "품절", "option_soldout": "옵션 품절", "moq": "최소수량 변경", "removed": "페이지 삭제", "restored": "복구(재판매)"}


def event_text(ev: dict, listing: dict, item: dict) -> str:
    a, b = ev.get("after_json") or {}, ev.get("before_json") or {}
    head = f"[{KIND_KO.get(ev['kind'], ev['kind'])}] {item.get('title_raw', '')[:40]}"
    if ev["kind"] == "price":
        body = f"단가 ¥{b.get('unit_price')} → ¥{a.get('unit_price')} ({a.get('change_pct'):+.0f}%)"
    elif ev["kind"] == "moq":
        body = f"최소수량 {b.get('moq')} → {a.get('moq')}"
    elif ev["kind"] == "option_soldout":
        body = "품절 옵션: " + ", ".join(a.get("newly") or [])
    else:
        body = KIND_KO.get(ev["kind"], "")
    tail = []
    if ev.get("recalculated_margin") is not None:
        tail.append(f"현재 판매가 {listing.get('sale_price_krw', 0):,}원 기준 재계산 마진율 {ev['recalculated_margin'] * 100:.0f}%")
    if ev.get("suggested_price_krw"):
        tail.append(f"목표 마진 판매가 제안 {ev['suggested_price_krw']:,}원")
    return "\n".join([head, body] + tail + [f"스토어 #{listing.get('store_product_no') or '-'} · {item.get('url', '')}"])


def send_cards(c, bot, events: list[dict]) -> int:
    from app.adapters.telegram import callback_data

    n = 0
    for ev in events:
        listing = c.table("listings").select("*").eq("candidate_id", ev["candidate_id"]).limit(1).execute().data[0]
        item = c.table("source_items").select("title_raw, url").eq("id", listing["source_item_id"]).limit(1).execute().data[0]
        btn = [[("판매가 수정", callback_data("m", ev["id"], "reprice")), ("판매 중지", callback_data("m", ev["id"], "suspend")), ("무시", callback_data("m", ev["id"], "ignore"))]]
        bot.send_text(event_text(ev, listing, item), buttons=btn)
        c.table("monitor_events").update({"notified_at": _now()}).eq("id", ev["id"]).execute()
        n += 1
    return n


# ───────────────────────── 결정 ─────────────────────────
def _append_update_xlsx(name: str, row: dict) -> Path:
    from openpyxl import Workbook, load_workbook

    path = ROOT / "output" / date.today().isoformat() / name
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = list(row.keys())
    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(cols)
    ws.append([row[k] for k in cols])
    wb.save(path)
    return path


def decide_monitor(c, uid: str, event_id: int, action: str) -> str:
    ev = c.table("monitor_events").select("*").eq("id", event_id).limit(1).execute().data
    if not ev:
        return "이벤트 없음"
    ev = ev[0]
    listing = c.table("listings").select("*").eq("candidate_id", ev["candidate_id"]).limit(1).execute().data[0]
    now = _now()
    if action == "reprice":
        new_price = int(ev.get("suggested_price_krw") or listing["sale_price_krw"])
        _append_update_xlsx("naver_price_update.xlsx", {"스토어상품번호": listing.get("store_product_no") or "", "현재판매가": listing["sale_price_krw"], "새판매가": new_price, "후보ID": listing["candidate_id"], "사유": ev["kind"]})
        item = c.table("source_items").select("price_tiers_json, moq, sold_out").eq("id", listing["source_item_id"]).limit(1).execute().data[0]
        det = (c.table("source_item_details").select("options_json").eq("source_item_id", listing["source_item_id"]).limit(1).execute().data or [{}])[0]
        base = snapshot(item.get("price_tiers_json") or [], item.get("moq") or 1, det.get("options_json") or [], bool(item.get("sold_out")))
        c.table("listings").update({"sale_price_krw": new_price, "baseline_json": base, "monitor_status": "ok"}).eq("candidate_id", listing["candidate_id"]).execute()
        label = f"판매가 수정 → {new_price:,}원 (naver_price_update.xlsx)"
    elif action == "suspend":
        _append_update_xlsx("naver_status_update.xlsx", {"스토어상품번호": listing.get("store_product_no") or "", "상태": "판매중지", "후보ID": listing["candidate_id"], "사유": ev["kind"]})
        c.table("listings").update({"suspended_at": now}).eq("candidate_id", listing["candidate_id"]).execute()
        c.table("candidates").update({"status": "suspended"}).eq("id", listing["candidate_id"]).execute()
        label = "판매 중지 (naver_status_update.xlsx)"
    elif action == "ignore":
        label = "무시 (다음 주기에 다시 비교)"
    else:
        return "알 수 없는 동작"
    c.table("monitor_events").update({"decision": action, "decided_at": now}).eq("id", event_id).execute()
    return label


def run(c, uid: str, *, only: int | None = None, bot=None, headed: bool = False) -> dict:
    q = c.table("listings").select("*")
    if only:
        q = q.eq("candidate_id", only)
    listings = q.execute().data
    out = {"checked": 0, "events": 0, "errors": [], "sent": 0}
    for lst in listings:
        res = check_one(c, uid, lst, headed=headed)
        out["checked"] += 1
        errs = [r["error"] for r in res if "error" in r]
        evs = [r for r in res if "id" in r]
        out["errors"] += [f"#{lst['candidate_id']}: {e}" for e in errs]
        out["events"] += len(evs)
        if bot is not None and evs:
            out["sent"] += send_cards(c, bot, evs)
    if bot is not None and out["errors"]:
        bot.send_text(f"감시 실패 {len(out['errors'])}건: " + " / ".join(out["errors"])[:500])
    return out


def main(argv=None) -> int:
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8")
        except Exception:
            pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--cand", type=int)
    ap.add_argument("--mark", type=int, help="등록 완료 표시할 후보 id")
    ap.add_argument("--store-no")
    ap.add_argument("--price", type=int)
    ap.add_argument("--no-telegram", action="store_true")
    ap.add_argument("--headed", action="store_true")
    a = ap.parse_args(argv)
    from app.db import current_user_id, user_client
    from app.user_settings import ensure_user_settings

    c = user_client()
    uid = current_user_id(c)
    ensure_user_settings(c, uid)
    if a.mark:
        cand = c.table("candidates").select("suggested_price_krw").eq("id", a.mark).limit(1).execute().data[0]
        row = mark_listed(c, uid, a.mark, store_product_no=a.store_no, sale_price_krw=a.price or cand["suggested_price_krw"] or 0)
        print("listed:", row["candidate_id"], "판매가", row["sale_price_krw"], "기준", row["baseline_json"])
        return 0
    bot = None
    if not a.no_telegram:
        try:
            from app.adapters.telegram import bot_from

            bot = bot_from(c)
        except Exception:
            pass
    print(run(c, uid, only=a.cand, bot=bot, headed=a.headed))
    return 0


if __name__ == "__main__":
    sys.exit(main())
