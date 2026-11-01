"""⑨ 패키지 생성 (DESIGN.md 2장): 이미지 → Supabase Storage 공개 URL → 네이버 일괄등록 엑셀 행 + packages 행.

  python -m app.pipeline.package [--cand 17]      # approved_2 → packaged

엑셀: output/YYYY-MM-DD/naver_bulk_upload.xlsx 에 행 추가(없으면 헤더부터). 컬럼은 네이버 "상품 일괄등록" 양식을 받아 매핑을
확정하기 전까지의 작업용 구성이다 (README 7장 "열려 있는 것"). 이미지 경로: <user_id>/<날짜>/<소스>_<상품ID>/<n>.jpg (Storage 정책이 본인 폴더만 허용).
"""
from __future__ import annotations

import argparse
import sys
from datetime import date, datetime, timezone
from pathlib import Path

from app.config import ROOT, settings
from app.vision import phash

EXCEL_COLUMNS = ["상품명", "판매가", "재고수량", "카테고리ID", "대표이미지", "추가이미지", "상세설명", "옵션명", "옵션값", "옵션가", "원산지", "배송비", "A/S안내", "소스URL", "소스상품ID", "후보ID"]


def excel_row(cand: dict, content: dict, item: dict, image_urls: list[str], *, naver_cat_id: str | None, country: str) -> dict:
    opts = [o for o in (content.get("options_ko_json") or []) if o.get("sell", True)]
    option_values = ",".join(o["name_ko"] for o in opts)
    # 옵션가: 소스 옵션가 차이를 원화로 반영하지 않고(1차) 0 으로 둔다 — 대표 옵션 기준 판매가
    option_prices = ",".join("0" for _ in opts)
    return {
        "상품명": content.get("title_chosen") or (content.get("title_options") or [""])[0],
        "판매가": int(cand.get("suggested_price_krw") or 0),
        "재고수량": 99,
        "카테고리ID": naver_cat_id or "",
        "대표이미지": image_urls[0] if image_urls else "",
        "추가이미지": ",".join(image_urls[1:9]),
        "상세설명": content.get("description_html") or "",
        "옵션명": "옵션" if opts else "",
        "옵션값": option_values,
        "옵션가": option_prices,
        "원산지": {"CN": "중국", "JP": "일본"}.get(country, country),
        "배송비": 0,
        "A/S안내": "해외 구매대행 상품으로 A/S 는 판매자 문의를 통해 안내드립니다.",
        "소스URL": item.get("url", ""),
        "소스상품ID": item.get("source_item_id", ""),
        "후보ID": cand["id"],
    }


def append_excel(path: Path, row: dict) -> None:
    from openpyxl import Workbook, load_workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "상품"
        ws.append(EXCEL_COLUMNS)
    ws.append([row.get(col, "") for col in EXCEL_COLUMNS])
    wb.save(path)


def upload_images(c, uid: str, cand: dict, item: dict, urls: list[str], *, bucket: str | None = None) -> list[str]:
    bucket = bucket or settings.supabase_bucket
    day = date.today().isoformat()
    folder = f"{uid}/{day}/{item['source_id']}_{item['source_item_id']}"
    out = []
    for i, u in enumerate(urls[:10], 1):
        ib = phash.fetch(u, timeout=15.0)
        if not ib:
            continue
        ext = "png" if "png" in (ib.content_type or "").lower() or u.lower().endswith(".png") else "jpg"
        path = f"{folder}/{i}.{ext}"
        try:
            c.storage.from_(bucket).upload(path, ib.data, {"content-type": ib.content_type or "image/jpeg", "upsert": "true"})
            out.append(c.storage.from_(bucket).get_public_url(path))
        except Exception as e:  # 한 장 실패가 전체를 막지 않게
            print(f"  업로드 실패 {path}: {type(e).__name__}: {str(e)[:120]}")
    return out


def build(c, uid: str, cand: dict) -> dict:
    content = c.table("generated_contents").select("*").eq("candidate_id", cand["id"]).limit(1).execute().data
    if not content:
        raise RuntimeError("생성 콘텐츠가 없습니다 (⑦ 먼저)")
    content = content[0]
    item = c.table("source_items").select("*").eq("id", cand["source_item_id"]).limit(1).execute().data[0]
    norm = (c.table("source_items_normalized").select("canonical_category_id").eq("source_item_id", item["id"]).limit(1).execute().data or [{}])[0]
    sig = (c.table("market_signals").select("naver_cat_id").eq("source_item_id", item["id"]).limit(1).execute().data or [{}])[0]
    naver_cat = sig.get("naver_cat_id")
    if not naver_cat and norm.get("canonical_category_id"):
        cc = c.table("canonical_categories").select("naver_cat_id").eq("id", norm["canonical_category_id"]).limit(1).execute().data
        naver_cat = (cc[0].get("naver_cat_id") if cc else None)
    src = (c.table("sources").select("country").eq("id", item["source_id"]).limit(1).execute().data or [{"country": "CN"}])[0]
    public_urls = upload_images(c, uid, cand, item, content.get("image_urls") or [])
    row = excel_row(cand, content, item, public_urls, naver_cat_id=naver_cat, country=src["country"])
    out_dir = ROOT / "output" / date.today().isoformat()
    xlsx = out_dir / "naver_bulk_upload.xlsx"
    append_excel(xlsx, row)
    c.table("packages").upsert({"candidate_id": cand["id"], "user_id": uid, "excel_row_json": row, "image_public_urls": public_urls, "output_path": str(xlsx.relative_to(ROOT)),
                                "sale_price_krw": row["판매가"]}, on_conflict="candidate_id").execute()
    c.table("candidates").update({"status": "packaged"}).eq("id", cand["id"]).execute()
    return {"xlsx": str(xlsx), "images": len(public_urls), "row": row}


def run(c, uid: str, *, only: int | None = None) -> dict:
    q = c.table("candidates").select("*").eq("status", "approved_2") if not only else c.table("candidates").select("*").eq("id", only)
    out = {"done": 0, "errors": [], "xlsx": None}
    for cand in q.execute().data:
        try:
            r = build(c, uid, cand)
            out["done"] += 1
            out["xlsx"] = r["xlsx"]
        except Exception as e:
            out["errors"].append(f"#{cand['id']}: {type(e).__name__}: {e}")
    return out


def main(argv=None) -> int:
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8")
        except Exception:
            pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--cand", type=int)
    a = ap.parse_args(argv)
    from app.db import current_user_id, user_client

    c = user_client()
    uid = current_user_id(c)
    r = run(c, uid, only=a.cand)
    print(r)
    return 0 if not r["errors"] else 1


if __name__ == "__main__":
    sys.exit(main())
